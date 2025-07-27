import docx
import streamlit as st
from fpdf import FPDF
from langchain_core.output_parsers import JsonOutputParser
from langchain_core.prompts import ChatPromptTemplate
from langchain_google_genai import ChatGoogleGenerativeAI
from pydantic import BaseModel, Field
from pypdf import PdfReader


# Configure page
st.set_page_config(
	page_title="TrackWise - ATS Resume Analyzer",
	page_icon="🎯",
	layout="wide",
	initial_sidebar_state="collapsed"
)

# Check for API key
if "GEMINI_API_KEY" not in st.secrets:
	st.error("🚨 Gemini API key not found! Please add it to your `.streamlit/secrets.toml` file.")
	st.stop()

st.html('style.css')


class ATSAnalysis(BaseModel):
	match_score: int = Field(description="integer between 0-100")
	summary: str = Field(description="brief 2-3 sentence summary of the candidate's fit")
	strengths: list[str] = Field(description="list of strengths")
	weaknesses: list[str] = Field(description="list of weaknesses")
	missing_keywords: list[str] = Field(description="list of missing keywords")
	improvement_suggestions: list[str] = Field(description="list of specific actionable suggestions")
	technical_skills_match: str = Field(description="percentage match for technical skills")
	experience_relevance: str = Field(description="assessment of experience relevance")
	education_match: str = Field(description="assessment of educational background match")
	recommendations: str = Field(description="specific recommendations for the candidate")
	interview_questions: list[dict] = Field(description="list of potential interview questions with answers based on resume")


def extract_text_from_file(uploaded_file) -> str:
	"""Extract text from the uploaded PDF or DOCX file"""
	try:
		file_type = uploaded_file.name.split('.')[-1].lower()
		
		if file_type == 'pdf':
			# Extract text from PDF
			pdf_reader = PdfReader(uploaded_file)
			text = ""
			for page in pdf_reader.pages:
				text += page.extract_text() + "\n"
			return text.strip()
			
		elif file_type == 'docx':
			# Extract text from DOCX
			doc = docx.Document(uploaded_file)
			text = "\n".join([paragraph.text for paragraph in doc.paragraphs])
			return text.strip()
			
		else:
			st.error(f"Unsupported file format: {file_type}. Please upload a PDF or DOCX file.")
			return ""
			
	except Exception as e:
		st.error(f"Error extracting text from file: {str(e)}")
		return ""

@st.cache_data	# DELETE  THIS BEFORE COMMITTING 
def analyze_resume(job_role: str, job_description: str, resume_text: str) -> ATSAnalysis | None:
	"""Analyze resume using LangChain and Google Gemini"""
	try:
		# Set up the parser
		parser = JsonOutputParser(pydantic_object=ATSAnalysis)
		
		# Create the prompt template
		with open('prompt.md') as f:
			prompt_template = ChatPromptTemplate.from_template(
				template=f.read(),
				partial_variables={"format_instructions": parser.get_format_instructions()}
			)
		
		# Initialize the model
		model = ChatGoogleGenerativeAI(
			model="gemini-2.5-flash",
			google_api_key=st.secrets["GEMINI_API_KEY"],
			temperature=0, thinking_budget=0
		)
		
		# Create the chain
		chain = prompt_template | model | parser
		
		# Invoke the chain
		response = chain.invoke({
			"job_role": job_role,
			"job_description": job_description,
			"resume_text": resume_text
		})
		
		return ATSAnalysis(**response)
	
	except Exception as e:
		st.error(f"Error analyzing resume: {e}")
		return None


def create_qna_pdf(interview_qna: list[dict]) -> bytes:
	"""Create a PDF file for the interview Q&A"""
	pdf = FPDF()
	pdf.add_page()
	pdf.set_font("Arial", 'B', 16)
	pdf.cell(0, 12, 'Interview Questions and Answers', 0, 1, 'C')
	pdf.ln(10)
	
	pdf.set_font("Arial", '', 12)
	for i, qa in enumerate(interview_qna, 1):
		pdf.set_font("Arial", 'B', 12)
		pdf.multi_cell(0, 7, f"{i}: {qa['question']}".encode('latin-1', 'replace').decode('latin-1'))
		pdf.set_font("Arial", '', 12)
		pdf.multi_cell(0, 7, f"{qa['answer']}".encode('latin-1', 'replace').decode('latin-1'))
		pdf.ln(5)
	return pdf.output(dest='S').encode('latin-1') # type: ignore


def display_results(analysis: ATSAnalysis, mode: str):
	"""Display analysis results with enhanced UI"""
	
	if mode == 'analysis':
		# Score Display
		st.html(f"""
		<div class="score-container">
			<div class="score-number">{analysis.match_score}%</div>
			<div class="score-label">ATS Match Score</div>
		</div>
		""")
		
		# Summary
		st.markdown("### 📋 Executive Summary")
		st.info(analysis.summary)
		
		# Create two columns for strengths and weaknesses
		col1, col2 = st.columns(2)
		
		with col1:
			st.markdown("### ✅ Strengths")
			for strength in analysis.strengths:
				st.markdown(f"• {strength}")
		
		with col2:
			st.markdown("### ⚠️ Areas for Improvement")
			for weakness in analysis.weaknesses:
				st.markdown(f"• {weakness}")
		
		# Missing Keywords
		if analysis.missing_keywords:
			st.markdown("### 🔍 Missing Keywords")
			keywords_str = ", ".join(analysis.missing_keywords)
			st.warning(f"Consider adding these keywords: **{keywords_str}**")
		
		# Detailed Analysis
		st.markdown("### 📊 Detailed Analysis")
		
		col3, col4, col5 = st.columns(3)
		with col3:
			st.write(f'**Technical Skills Match:** \n\n {analysis.technical_skills_match}')
		with col4:
			st.write(f'**Experience Relevance:** \n\n {analysis.experience_relevance}')
		with col5:
			st.write(f'**Education Match:** \n\n {analysis.education_match}')
		
		# Improvement Suggestions
		st.html(f"""
		<div class="improvement-section">
			<h3>🚀 Improvement Suggestions</h3>
			{''.join([f"<p>{i}. {suggestion}</p>" for i, suggestion in enumerate(analysis.improvement_suggestions, 1)])}
		</div>
		""")

		# Recommendations
		if analysis.recommendations:
			st.html(f"""
			<div class="recommendation-section">
				<h3>💡 Professional Recommendations</h3>
				<p>{analysis.recommendations}</p>
			</div>
			""")
	
	elif mode == 'interview':
		# Interview Questions
		if analysis.interview_questions:
			st.markdown("### 🎯 Prepare for Your Interview")
			st.markdown("Here are some potential interview questions based on your resume.")


			col1, col2 = st.columns(2)
			
			for index, qa in enumerate(analysis.interview_questions):
				col = col1 if index % 2 == 0 else col2
				with col:
					st.html(f"""
						<div class="qa-card">
							<p class="question">{qa['question']}</p>
							<p style="color: #333333;">{qa['answer']}</p>
						</div>
					""")

			# Add download button
			pdf_bytes = create_qna_pdf(analysis.interview_questions)
			st.download_button(
				label="Export as PDF",
				icon=":material/file_save:",
				data=pdf_bytes,
				file_name="interview_qna.pdf",
				type="primary",
				on_click="ignore"
			)

# Header
st.html("""
<div class="main-header">
	<h1>🎯 TrackWise</h1>
	<p>Advanced ATS Resume Analyzer &nbsp;•&nbsp; Optimize Your Resume for Success</p>
</div>
""")

# Input fields
col1, col2 = st.columns([1, 1])

with col1:
	job_role = st.text_input("🎯 Job Role", placeholder="e.g., Data Analyst, Software Engineer")
	uploaded_file = st.file_uploader("📄 Upload Resume", type=['pdf', 'docx'], help="Upload your resume in PDF or DOCX format")

with col2:
	job_description = st.text_area(
		"📝 Job Description",
		placeholder="Paste the complete job description here...",
		height=300
	)

# Action buttons
_, col1, col2, _ = st.columns([1, 2, 2, 1], gap="medium")
with col1:
	analyze_button = st.button("Analyze Resume", use_container_width=True)
with col2:
	interview_button = st.button("Prepare for Interview", use_container_width=True)

if 'analysis_result' not in st.session_state:
	st.session_state.analysis_result = None

def run_analysis(mode: str):
	if not job_role.strip():
		st.error("⚠️ Please enter the job role")
	elif not job_description.strip():
		st.error("⚠️ Please enter the job description")
	elif not uploaded_file:
		st.error("⚠️ Please upload your resume")
	else:
		spinner_text = "Analyzing your resume..." if mode == 'analysis' else "Preparing interview questions..."
		with st.spinner(spinner_text):
			resume_text = extract_text_from_file(uploaded_file)
			if resume_text:
				st.session_state.analysis_result = analyze_resume(job_role, job_description, resume_text)
				if not st.session_state.analysis_result:
					st.error("❌ Failed to analyze the resume. Please try again.")
			else:
				st.error("❌ Failed to extract text from file. Please ensure the file is readable and in the correct format.")

if analyze_button:
	run_analysis('analysis')
	if st.session_state.analysis_result:
		display_results(st.session_state.analysis_result, 'analysis')

if interview_button:
	run_analysis('interview')
	if st.session_state.analysis_result:
		display_results(st.session_state.analysis_result, 'interview')

# Footer
st.divider()
st.html("""
<div style="text-align: center; color: #666;">
	<p style="font-size: 0.95rem;">Built by <a href="https://subhankar-dutta.streamlit.app" target="_blank"><b>Subhankar Dutta</b></a>
	 | Powered by <span class="gemini-text"><b>Gemini</b></span></p>
	<p style="font-size: 0.8rem;">Upload your resume and get instant ATS optimization insights</p>
</div>
""")